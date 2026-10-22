"""Render the 48 x 36 in poster to PDF (print and Canva import), PowerPoint and a preview PNG.

Needs the dashboard dev server: pnpm --dir dashboard dev --port 3123
Usage: python scripts/render_poster.py dashboard/public/poster/exports [--phone]
--phone first replays one /drive round and saves the phone screenshot the poster embeds.
Prints the smallest text size in points and any region whose content overflows.

The PowerPoint is one 48 x 36 in slide: the poster captured at 144 dpi with its HTML text made
transparent, under native text boxes holding that text in place. Text drawn inside SVG figures stays
in the picture. The boxes name IBM Plex Sans, which PowerPoint substitutes where it is not installed.
"""
import asyncio
import re
import sys
from pathlib import Path

from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

BASE = "http://localhost:3123"
REPO = Path(__file__).resolve().parents[1]
HIDE_DEV = "nextjs-portal{display:none!important}[data-screen-only]{display:none!important}"
POSTER_PX = (4608, 3456)
EMU_PER_PX = 914400 // 96
PT_PER_PX = 0.75
# The layout loads IBM's static TTFs, whose family names carry the weight.
PLEX_BY_WEIGHT = {400: ("IBM Plex Sans", False), 500: ("IBM Plex Sans Medium", False),
                  600: ("IBM Plex Sans SemiBold", False), 700: ("IBM Plex Sans", True)}
ALIGN = {"right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}

# Every HTML text block outside SVG figures, as positioned runs; each is then marked so CSS can blank it.
TEXT_BOXES = """() => {
  const main = document.querySelector('main');
  const origin = main.getBoundingClientRect();
  const blocks = [...main.querySelectorAll('h1, h2, h3, p, figcaption')].filter(el =>
    !el.closest('svg') && el.textContent.trim() && !el.querySelector('h1, h2, h3, p, figcaption'));
  return blocks.map(el => {
    const runs = [];
    const add = (text, styled) => {
      const s = getComputedStyle(styled);
      runs.push({ text, size: parseFloat(s.fontSize), weight: parseInt(s.fontWeight, 10), color: s.color,
                  lineHeight: parseFloat(s.lineHeight), sup: !!styled.closest('sup') });
    };
    // A child starts a new line when it is a block, or an item of a column flexbox or a grid.
    const breaksLine = (parent, child) => {
      const p = getComputedStyle(parent), c = getComputedStyle(child);
      const stacked = (p.display.includes('flex') && p.flexDirection.startsWith('column')) || p.display.includes('grid');
      return stacked || c.display === 'block';
    };
    const walk = node => {
      for (const child of node.childNodes) {
        if (child.nodeType === Node.TEXT_NODE) {
          const text = child.textContent.replace(/[ \\t\\n\\r]+/g, ' ');
          if (text.trim() || (text && runs.length)) add(text, child.parentElement);
        } else if (child.nodeType === Node.ELEMENT_NODE) {
          if (child.tagName === 'BR' || (runs.length && breaksLine(node, child))) add('\\n', node);
          else if (runs.length && parseFloat(getComputedStyle(child).marginLeft) > 0) add('   ', child);
          if (child.tagName !== 'BR') walk(child);
        }
      }
    };
    walk(el);
    const box = el.getBoundingClientRect();
    const style = getComputedStyle(el);
    el.setAttribute('data-pptx-text', '');
    return { left: box.left - origin.left, top: box.top - origin.top, width: box.width, height: box.height,
             align: style.textAlign, runs };
  });
}"""
BLANK_TEXT = "[data-pptx-text], [data-pptx-text] * { color: transparent !important; }"

async def phone_shot(b):
    ctx = await b.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=3, is_mobile=True, has_touch=True)
    pg = await ctx.new_page()
    await pg.goto(f"{BASE}/drive", wait_until="networkidle")
    await pg.add_style_tag(content=HIDE_DEV)
    # The top-down view exposes the pedestrian as SVG, so it times the brake; the poster shows the 3D result.
    await pg.get_by_role("radio", name="Top-down").click()
    await pg.get_by_role("button", name="Start").click()
    await pg.wait_for_function("document.querySelectorAll('svg ellipse').length >= 3", timeout=10000)
    await pg.wait_for_timeout(200)
    box = await pg.get_by_role("button", name="Brake").bounding_box()
    await pg.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
    await pg.mouse.down()
    await pg.wait_for_selector("text=Try again", timeout=10000)
    await pg.mouse.up()
    await pg.get_by_role("radio", name="3D street").click()
    await pg.wait_for_timeout(2000)
    await pg.evaluate("window.scrollTo(0,0)")
    await pg.screenshot(path=f"{REPO}/dashboard/public/poster/drive-phone.png")
    await ctx.close()

AUDIT = """() => {
  let minPt = 1e9, smallest = [];
  for (const el of document.querySelectorAll('main p, main h1, main h2, main h3, main text, main tspan, main span, main figcaption, main sup')) {
    const ownText = [...el.childNodes].some(n => n.nodeType === 3 && n.textContent.trim());
    if (!ownText) continue;
    const fs = parseFloat(getComputedStyle(el).fontSize);
    let scale = 1;
    const svg = el.closest('svg');
    if (svg && svg.viewBox.baseVal && svg.viewBox.baseVal.width) scale = svg.getBoundingClientRect().width / svg.viewBox.baseVal.width;
    const pt = Math.round(fs * scale * 0.75 * 10) / 10;
    if (pt < minPt) { minPt = pt; smallest = []; }
    if (pt === minPt) smallest.push(el.textContent.trim().slice(0, 30));
  }
  const overflow = [];
  for (const region of document.querySelectorAll('main > div')) {
    const r = region.getBoundingClientRect();
    for (const child of region.querySelectorAll('*:not(svg *)')) {
      const c = child.getBoundingClientRect();
      if (c.width && (c.bottom > r.bottom + 1 || c.right > r.right + 1)) { overflow.push((region.querySelector('h2,h1')||{}).textContent + ' <- ' + child.tagName + ' ' + Math.round(c.bottom - r.bottom)); break; }
    }
  }
  return { minPt, smallest: smallest.slice(0, 5), overflow };
}"""

async def poster(b, out):
    ctx = await b.new_context(viewport={"width": 4608, "height": 3456})
    pg = await ctx.new_page()
    await pg.goto(f"{BASE}/poster", wait_until="networkidle")
    await pg.add_style_tag(content=HIDE_DEV)
    await pg.evaluate("document.fonts.ready")
    audit = await pg.evaluate(AUDIT)
    print("min text pt:", audit["minPt"], audit["smallest"])
    print("overflowing regions:", audit["overflow"])
    await pg.screenshot(path=f"{out}/poster-preview.png")
    await pg.pdf(path=f"{out}/poster.pdf", width="48in", height="36in", print_background=True, margin={"top": "0", "right": "0", "bottom": "0", "left": "0"})
    await ctx.close()

def px(value):
    return Emu(round(value * EMU_PER_PX))


def rgb(css_color):
    red, green, blue = (int(float(channel)) for channel in re.findall(r"[\d.]+", css_color)[:3])
    return RGBColor(red, green, blue)


def style_run(run, spec):
    family, bold = PLEX_BY_WEIGHT.get(spec["weight"], PLEX_BY_WEIGHT[400])
    run.text = spec["text"]
    run.font.name = family
    run.font.bold = bold
    run.font.size = Pt(spec["size"] * PT_PER_PX)
    run.font.color.rgb = rgb(spec["color"])
    if spec["sup"]:
        run.font._rPr.set("baseline", "30000")


def paragraphs(runs):
    """Split runs at line breaks into one list of runs per paragraph, trimmed at both ends."""
    lines = [[]]
    for spec in runs:
        if spec["text"] == "\n":
            lines.append([])
        else:
            lines[-1].append(dict(spec))
    lines = [line for line in lines if any(spec["text"].strip() for spec in line)]
    for line in lines:
        line[0]["text"] = line[0]["text"].lstrip()
        line[-1]["text"] = line[-1]["text"].rstrip()
    return lines


def add_text_box(slide, block):
    frame = slide.shapes.add_textbox(px(block["left"]), px(block["top"]), px(block["width"] * 1.01),
                                     px(block["height"])).text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(paragraphs(block["runs"])):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = ALIGN.get(block["align"], PP_ALIGN.LEFT)
        paragraph.line_spacing = Pt(line[0]["lineHeight"] * PT_PER_PX)
        for spec in line:
            style_run(paragraph.add_run(), spec)


def write_pptx(path, background, blocks):
    deck = Presentation()
    deck.slide_width, deck.slide_height = px(POSTER_PX[0]), px(POSTER_PX[1])
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(background), 0, 0, deck.slide_width, deck.slide_height)
    for block in blocks:
        add_text_box(slide, block)
    deck.save(path)


async def powerpoint(b, out):
    # 1.5x (144 dpi) is the largest capture Chromium accepts here; 2x failed. The PDF is the print file.
    ctx = await b.new_context(viewport={"width": POSTER_PX[0], "height": POSTER_PX[1]}, device_scale_factor=1.5)
    pg = await ctx.new_page()
    await pg.goto(f"{BASE}/poster", wait_until="networkidle")
    await pg.add_style_tag(content=HIDE_DEV)
    await pg.evaluate("document.fonts.ready")
    blocks = await pg.evaluate(TEXT_BOXES)
    await pg.add_style_tag(content=BLANK_TEXT)
    background = Path(out) / "poster-background.jpg"
    await pg.locator("main").screenshot(path=str(background), type="jpeg", quality=90)
    await ctx.close()
    write_pptx(f"{out}/poster.pptx", background, blocks)
    background.unlink()
    print("powerpoint text boxes:", len(blocks))


async def main():
    out = sys.argv[1]
    async with async_playwright() as p:
        b = await p.chromium.launch(args=["--use-angle=swiftshader", "--enable-unsafe-swiftshader"])
        if "--phone" in sys.argv:
            await phone_shot(b)
        await poster(b, out)
        await powerpoint(b, out)
        await b.close()
asyncio.run(main())
